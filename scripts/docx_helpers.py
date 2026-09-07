"""
ThermoTrace Document Generator
Generates:
1. docs/THERMOTRACE_END_TO_END_IMPLEMENTATION_MASTER.docx
2. docs/THERMOTRACE_SIH26162_EVALUATION_CRITERIA_DEFENSE.docx
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    if level == 1:
        run.font.name = 'Arial'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(11, 23, 40) # Deep Navy
        pPr = h._element.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="008080"/></w:pBdr>')
        pPr.append(pBdr)
    elif level == 2:
        run.font.name = 'Arial'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 102, 153) # Steel Blue
    elif level == 3:
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(40, 50, 60)
    return h

def add_callout(doc, title, text, border_color="008080", fill_color="F0F8FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(0, 70, 120)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def format_table_header(table, cols):
    for j, c in enumerate(cols):
        cell = table.cell(0, j)
        set_cell_background(cell, "0B1728")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.runs[0] if p.runs else p.add_run(c)
        r.text = c
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

def format_table_rows(table, rows_data):
    for i, row in enumerate(rows_data):
        row_idx = i + 1
        bg_col = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row):
            cell = table.cell(row_idx, j)
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r = p.runs[0] if p.runs else p.add_run(str(val))
            r.text = str(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 30, 30)

print("Setup completed. Helper functions ready.")
